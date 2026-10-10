"""Append-only, research-only forward snapshots of the official China EPU file."""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import math
import os
import ssl
import tempfile
import urllib.request
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

from . import config


SCOUT_ID = "bcc5cc436207a3e589be"
SOURCE_URL = "https://www.policyuncertainty.com/media/China_Mainland_Paper_EPU.xlsx"
SHEET = "EPU 2000 onwards"
OUTPUT_DIR = config.PROJECT_ROOT / "research/data/epu_beta_v1_epu_forward_snapshots"
MAX_BYTES = 10 * 1024 * 1024


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _relative(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(config.PROJECT_ROOT.resolve()))
    except ValueError:
        return str(path.resolve())


def _write_once(path: Path, data: bytes) -> None:
    """Atomically create a content-addressed artifact without replacing bytes."""
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(dir=path.parent, prefix=".snapshot-")
    try:
        with os.fdopen(descriptor, "wb") as target:
            target.write(data)
            target.flush()
            os.fsync(target.fileno())
        try:
            os.link(temporary, path)
        except FileExistsError:
            if path.read_bytes() != data:
                raise ValueError(f"append-only snapshot collision: {path.name}")
    finally:
        Path(temporary).unlink(missing_ok=True)


def _default_fetch() -> tuple[bytes, dict[str, str | None]]:
    request = urllib.request.Request(
        SOURCE_URL,
        headers={"User-Agent": "RegisOracle-EPU-PIT-Snapshot/1.0"},
    )
    paths = ssl.get_default_verify_paths()
    cafile = next((candidate for candidate in (
        paths.cafile, paths.openssl_cafile, "/etc/ssl/cert.pem",
    ) if candidate and Path(candidate).is_file()), None)
    context = ssl.create_default_context(cafile=cafile)
    with urllib.request.urlopen(request, timeout=20, context=context) as response:
        data = response.read(MAX_BYTES + 1)
        metadata = {
            "content_type": response.headers.get("Content-Type"),
            "etag": response.headers.get("ETag"),
            "last_modified": response.headers.get("Last-Modified"),
        }
    if len(data) > MAX_BYTES:
        raise ValueError("official EPU workbook exceeds the 10 MiB safety limit")
    return data, metadata


def _inspect(data: bytes) -> dict:
    if not data.startswith(b"PK"):
        raise ValueError("official EPU response is not an XLSX workbook")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise ValueError(f"official EPU workbook is missing sheet: {SHEET}")
    rows = list(workbook[SHEET].iter_rows(min_col=1, max_col=3, values_only=True))
    if not rows or tuple(str(value).strip().casefold() for value in rows[0]) != (
        "year", "month", "epu",
    ):
        raise ValueError("official EPU workbook header changed")
    months: list[str] = []
    values: list[float] = []
    for year, month, value in rows[1:]:
        if year is None and month is None and value is None:
            continue
        try:
            year_int, month_int, number = int(year), int(month), float(value)
        except (TypeError, ValueError) as exc:
            raise ValueError("official EPU workbook contains an invalid row") from exc
        if not 1 <= month_int <= 12 or not math.isfinite(number) or number < 0:
            raise ValueError("official EPU workbook contains an invalid month or value")
        months.append(f"{year_int:04d}-{month_int:02d}")
        values.append(number)
    if not months or len(months) != len(set(months)):
        raise ValueError("official EPU workbook months are empty or duplicated")
    expected = []
    cursor = dt.date.fromisoformat(months[0] + "-01")
    finish = dt.date.fromisoformat(months[-1] + "-01")
    while cursor <= finish:
        expected.append(cursor.strftime("%Y-%m"))
        cursor = (cursor.replace(day=28) + dt.timedelta(days=4)).replace(day=1)
    if months != expected:
        raise ValueError("official EPU workbook monthly sequence is not contiguous")
    return {
        "sheet": SHEET,
        "first_month": months[0],
        "last_month": months[-1],
        "monthly_rows": len(months),
        "values_finite_nonnegative": len(values) == len(months),
    }


def capture(
    *, now: dt.datetime | None = None,
    output_dir: Path | None = None,
    fetcher: Callable[[], bytes | tuple[bytes, dict]] | None = None,
) -> dict:
    """Capture one observation; never compute a factor, rank, return, or signal."""
    now = now or dt.datetime.now().astimezone()
    output_dir = output_dir or OUTPUT_DIR
    if now.tzinfo is None:
        raise ValueError("EPU snapshot time must be timezone-aware")
    fetched = (fetcher or _default_fetch)()
    data, headers = fetched if isinstance(fetched, tuple) else (fetched, {})
    if not isinstance(data, bytes) or not data:
        raise ValueError("official EPU snapshot is empty")
    if len(data) > MAX_BYTES:
        raise ValueError("official EPU workbook exceeds the 10 MiB safety limit")
    workbook = _inspect(data)
    digest = _sha256(data)
    raw_path = output_dir / "raw" / f"{digest}.xlsx"
    _write_once(raw_path, data)
    stamp = now.isoformat(timespec="seconds")
    observation = {
        "schema_version": 1,
        "status": "accepted_forward_input_snapshot_not_factor",
        "candidate_id": "epu_beta_v1",
        "scout_id": SCOUT_ID,
        "source_url": SOURCE_URL,
        "fetched_at": stamp,
        "source_sha256": digest,
        "source_bytes": len(data),
        "raw_path": _relative(raw_path),
        "http": {
            key: headers.get(key) for key in ("content_type", "etag", "last_modified")
        },
        "workbook": workbook,
        "license": "free use with attribution to authors, paper, and policyuncertainty.com",
        "audit_boundary": {
            "historical_vintage_inferred": False,
            "publication_lag_inferred": False,
            "factor_values_computed": False,
            "returns_read": False,
            "ic_or_backtest_computed": False,
            "recommendation_or_production_allowed": False,
        },
    }
    payload = (json.dumps(observation, ensure_ascii=False, indent=2) + "\n").encode()
    name = now.strftime("%Y%m%dT%H%M%S%z") + f"-{digest[:12]}.json"
    observation_path = output_dir / "observations" / name
    _write_once(observation_path, payload)
    return {
        "status": observation["status"],
        "fetched_at": stamp,
        "source_sha256": digest,
        "last_month": workbook["last_month"],
        "monthly_rows": workbook["monthly_rows"],
        "observation_path": _relative(observation_path),
        "research_only": True,
        "runnable": False,
        "production_eligible": False,
    }


def capture_if_active(candidates: list[dict], **kwargs) -> dict:
    active = any(
        item.get("scout_id") == SCOUT_ID
        and item.get("status") == "source_contract_pending"
        and item.get("runnable") is False
        for item in candidates if isinstance(item, dict)
    )
    return capture(**kwargs) if active else {
        "status": "skipped_candidate_not_active",
        "research_only": True,
        "runnable": False,
        "production_eligible": False,
    }
