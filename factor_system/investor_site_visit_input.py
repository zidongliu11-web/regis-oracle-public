"""Outcome-blind source audit for abnormal investor site visits (AIV).

Eastmoney is used only to discover structured event rows.  An event is not
point-in-time eligible until its date, physical-visit marker and body content
bind to a CNINFO/互动易 PDF with an official publication timestamp.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
from pathlib import Path
from typing import Any, Callable, Iterable
from zoneinfo import ZoneInfo

import requests
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
AUDIT_PATH = ROOT / "research/INVESTOR_SITE_VISIT_AIV_V1_SOURCE_INPUT_AUDIT.json"
EASTMONEY_URL = "https://datacenter-web.eastmoney.com/api/data/v1/get"
CNINFO_SEARCH_URL = "https://irm.cninfo.com.cn/newircs/index/search"
CNINFO_HISTORY_SEARCH_URL = "https://www.cninfo.com.cn/new/hisAnnouncement/query"
CNINFO_STOCK_INDEX_URL = "https://www.cninfo.com.cn/new/data/szse_stock.json"
CNINFO_STOCK_SEARCH_URL = "https://www.cninfo.com.cn/new/information/topSearch/query"
CNINFO_PDF_BASE = "https://static.cninfo.com.cn/"
SHANGHAI = ZoneInfo("Asia/Shanghai")
REMOTE_OR_NON_SITE_TERMS = (
    "电话", "视频", "网络", "线上", "邮件", "路演", "业绩说明会",
    "股东大会", "新闻发布会", "媒体采访",
)
PHYSICAL_TERMS = (
    "现场参观", "现场调研", "实地考察", "公司会议室", "总部", "厂区", "工厂",
)
EASTMONEY_COLUMNS = (
    "SECUCODE,SECURITY_CODE,SECURITY_NAME_ABBR,NOTICE_DATE,"
    "RECEIVE_START_DATE,RECEIVE_END_DATE,RECEIVE_TIME_EXPLAIN,"
    "RECEIVE_OBJECT,RECEIVE_PLACE,RECEIVE_WAY_EXPLAIN,INVESTIGATORS,"
    "RECEPTIONIST,ORG_TYPE,NUMBERNEW,EITIME,EUTIME,CONTENT,IS_SOURCE"
)
USER_AGENT = "Mozilla/5.0 (Regis Oracle outcome-blind source audit)"


class SourceContractError(RuntimeError):
    """The public source failed a bounded, reproducible input contract."""


def _canonical(value: object) -> str:
    return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "", str(value or ""))


def _day(value: object) -> dt.date:
    return dt.datetime.fromisoformat(str(value or "")).date()


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def normalize_discovery_event(row: dict[str, Any]) -> dict[str, Any] | None:
    """Return a conservative physical SZSE event, never a tradable signal."""
    code = str(row.get("SECURITY_CODE") or "").strip()
    secucode = str(row.get("SECUCODE") or "").strip().upper()
    way = str(row.get("RECEIVE_WAY_EXPLAIN") or "").strip()
    place = str(row.get("RECEIVE_PLACE") or "").strip()
    if (
        not re.fullmatch(r"\d{6}", code)
        or not secucode.endswith(".SZ")
        or str(row.get("IS_SOURCE") or "") != "1"
        or str(row.get("NUMBERNEW") or "") != "1"
        or "现场参观" not in way
        or any(term in way + place for term in REMOTE_OR_NON_SITE_TERMS)
    ):
        return None
    try:
        visit_start = _day(row.get("RECEIVE_START_DATE"))
        visit_end = _day(row.get("RECEIVE_END_DATE") or row.get("RECEIVE_START_DATE"))
        discovery_notice = _day(row.get("NOTICE_DATE"))
    except (TypeError, ValueError):
        return None
    if visit_end < visit_start or discovery_notice < visit_start:
        return None

    content = " ".join(str(row.get("CONTENT") or "").split())
    identity_fields = (
        code, visit_start.isoformat(), visit_end.isoformat(), way, place,
        str(row.get("RECEPTIONIST") or "").strip(), _sha256_bytes(content.encode()),
    )
    event_id = hashlib.sha256("\x1f".join(identity_fields).encode()).hexdigest()
    return {
        "event_id": event_id,
        "stock_code": code,
        "stock_name": str(row.get("SECURITY_NAME_ABBR") or "").strip(),
        "visit_start_date": visit_start.isoformat(),
        "visit_end_date": visit_end.isoformat(),
        "discovery_notice_date": discovery_notice.isoformat(),
        "visit_time_text": str(row.get("RECEIVE_TIME_EXPLAIN") or "").strip(),
        "visit_way": way,
        "visit_place": place or None,
        "participant_type": str(row.get("ORG_TYPE") or "").strip() or None,
        "content": content,
        "content_sha256": _sha256_bytes(content.encode()),
        "secondary_ingested_at": str(row.get("EITIME") or "").strip() or None,
        "official_published_at": None,
        "point_in_time_eligible": False,
        "source_role": "secondary_discovery_only",
    }


def deduplicate_events(rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    unique: dict[str, dict[str, Any]] = {}
    for row in rows:
        event = normalize_discovery_event(row)
        if event is not None:
            unique[event["event_id"]] = event
    return [unique[key] for key in sorted(unique)]


def disclosure_session_lag(
    visit_date: str, publication_date: str, trading_dates: Iterable[str],
) -> int:
    visit = _day(visit_date)
    publication = _day(publication_date)
    return sum(visit < _day(value) <= publication for value in trading_dates)


def _date_markers(value: str) -> set[str]:
    day = _day(value)
    return {
        _canonical(day.isoformat()),
        _canonical(f"{day.year}年{day.month}月{day.day}日"),
        _canonical(f"{day.year}年{day.month:02d}月{day.day:02d}日"),
        _canonical(f"{day.month}月{day.day}日"),
        _canonical(f"{day.month:02d}月{day.day:02d}日"),
    }


def _content_anchor_hits(content: str, official_text: str, width: int = 18) -> int:
    source = _canonical(content)
    target = _canonical(official_text)
    if len(source) < width or not target:
        return 0
    starts = sorted({
        0,
        max(0, len(source) // 4 - width // 2),
        max(0, len(source) // 2 - width // 2),
        max(0, 3 * len(source) // 4 - width // 2),
        max(0, len(source) - width),
    })
    return sum(source[start:start + width] in target for start in starts)


def assess_official_text(event: dict[str, Any], official_text: str) -> dict[str, Any]:
    compact = _canonical(official_text)
    date_match = any(
        marker and marker in compact
        for marker in _date_markers(event["visit_start_date"])
    )
    physical_match = any(_canonical(term) in compact for term in PHYSICAL_TERMS)
    place = _canonical(event.get("visit_place"))
    place_match = bool(len(place) >= 6 and place in compact)
    anchor_hits = _content_anchor_hits(str(event.get("content") or ""), official_text)
    return {
        "visit_date_match": date_match,
        "physical_marker_match": physical_match or place_match,
        "content_anchor_hits": anchor_hits,
        "event_body_match": anchor_hits >= 1,
        "valid": bool(date_match and (physical_match or place_match) and anchor_hits >= 1),
    }


def _request_json(
    request: Callable[..., Any], *args: Any, attempts: int = 3, **kwargs: Any,
) -> dict[str, Any]:
    error: Exception | None = None
    for attempt in range(attempts):
        try:
            response = request(*args, **kwargs)
            response.raise_for_status()
            payload = response.json()
            if isinstance(payload, dict):
                return payload
            raise SourceContractError("source response is not an object")
        except (requests.RequestException, ValueError, SourceContractError) as exc:
            error = exc
            if attempt + 1 < attempts:
                time.sleep(0.15 * (attempt + 1))
    raise SourceContractError(f"source request failed: {error}")


def fetch_discovery_page(
    session: requests.Session, year: int, page: int, *, page_size: int = 50,
) -> dict[str, Any]:
    if not 1 <= page_size <= 50:
        raise ValueError("Eastmoney page_size must be between 1 and 50")
    filters = (
        '(IS_SOURCE="1")(NUMBERNEW="1")(SECUCODE like "%.SZ")'
        '(RECEIVE_WAY_EXPLAIN like "%现场参观%")'
        f"(RECEIVE_START_DATE>='{year}-01-01')"
        f"(RECEIVE_START_DATE<='{year}-12-31 23:59:59')"
    )
    params = {
        "sortColumns": "NOTICE_DATE,RECEIVE_START_DATE,SECURITY_CODE,NUMBERNEW",
        "sortTypes": "-1,-1,1,-1",
        "pageSize": str(page_size),
        "pageNumber": str(page),
        "reportName": "RPT_ORG_SURVEY",
        "columns": EASTMONEY_COLUMNS,
        "source": "WEB",
        "client": "WEB",
        "filter": filters,
    }
    payload = _request_json(
        session.get, EASTMONEY_URL, params=params, timeout=30,
    )
    result = payload.get("result")
    if payload.get("success") is not True or not isinstance(result, dict):
        raise SourceContractError(
            f"Eastmoney discovery unavailable: {payload.get('message')}"
        )
    return result


def fetch_official_documents(
    session: requests.Session, code: str, publication_date: str,
) -> list[dict[str, Any]]:
    form = {
        "pageNo": "1",
        "pageSize": "30",
        # The documented stockCode form field is ignored by the live endpoint;
        # an exact six-digit keyword is the bounded working lookup.
        "keyWord": code,
        "beginDate": f"{publication_date} 00:00:00",
        "endDate": f"{publication_date} 23:59:59",
        "stockCode": "",
        "searchTypes": "4",
        "highLight": "false",
    }
    payload: dict[str, Any] = {}
    for attempt in range(3):
        payload = _request_json(
            session.post, CNINFO_SEARCH_URL, data=form, timeout=30,
        )
        if payload.get("results"):
            break
        if attempt < 2:
            time.sleep(0.2 * (attempt + 1))
    pages = int(payload.get("totalPage") or 0)
    if pages > 1:
        raise SourceContractError("CNINFO exact code/date lookup exceeded one page")
    return [
        row for row in payload.get("results") or []
        if isinstance(row, dict) and str(row.get("stockCode") or "") == code
    ]


def fetch_cninfo_stock_org_ids(session: requests.Session) -> dict[str, str]:
    """Return the official code-to-orgId index used by historical disclosure search."""
    payload = _request_json(session.get, CNINFO_STOCK_INDEX_URL, timeout=30)
    rows = payload.get("stockList")
    if not isinstance(rows, list):
        raise SourceContractError("CNINFO stock index has no stockList")
    return {
        str(row.get("code")): str(row.get("orgId"))
        for row in rows
        if isinstance(row, dict)
        and re.fullmatch(r"\d{6}", str(row.get("code") or ""))
        and str(row.get("orgId") or "")
    }


def fetch_cninfo_stock_org_id(
    session: requests.Session, code: str,
) -> str:
    """Resolve an orgId for a code absent from the current stock index."""
    payload = session.post(
        CNINFO_STOCK_SEARCH_URL,
        data={"keyWord": code, "maxNum": "10", "plate": "szsh"},
        timeout=30,
    )
    payload.raise_for_status()
    try:
        rows = payload.json()
    except ValueError as exc:
        raise SourceContractError("CNINFO stock search is not JSON") from exc
    matches = [
        row for row in rows
        if isinstance(row, dict)
        and str(row.get("code") or "") == code
        and str(row.get("orgId") or "")
    ] if isinstance(rows, list) else []
    org_ids = sorted({str(row["orgId"]) for row in matches})
    if len(org_ids) != 1:
        raise SourceContractError(
            f"CNINFO exact stock orgId lookup returned {len(org_ids)} matches"
        )
    return org_ids[0]


def fetch_historical_official_documents(
    session: requests.Session,
    code: str,
    org_id: str,
    publication_date: str,
) -> list[dict[str, Any]]:
    """Query CNINFO's official historical investor-relations archive."""
    form = {
        "pageNum": "1",
        "pageSize": "30",
        "column": "szse",
        "tabName": "relation",
        "plate": "",
        "stock": f"{code},{org_id}",
        "searchkey": "",
        "secid": "",
        "category": "",
        "trade": "",
        "seDate": f"{publication_date}~{publication_date}",
        "sortName": "",
        "sortType": "",
        "isHLtitle": "true",
    }
    payload = _request_json(
        session.post, CNINFO_HISTORY_SEARCH_URL, data=form, timeout=30,
    )
    total = int(payload.get("totalRecordNum") or 0)
    if total > 30:
        raise SourceContractError("CNINFO exact code/date archive lookup exceeded one page")
    expected_day = _day(publication_date)
    documents = []
    for row in payload.get("announcements") or []:
        if not isinstance(row, dict) or str(row.get("secCode") or "") != code:
            continue
        try:
            actual_day = dt.datetime.fromtimestamp(
                int(str(row.get("announcementTime"))) / 1000, SHANGHAI,
            ).date()
        except (OSError, TypeError, ValueError):
            continue
        if actual_day == expected_day:
            documents.append(row)
    return documents


def _extract_original_attachment_text(raw: bytes, attachment_type: str) -> str:
    kind = attachment_type.upper()
    if kind == "PDF":
        if not raw.startswith(b"%PDF"):
            raise SourceContractError("official PDF signature is invalid")
        return "\n".join(
            page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages
        )
    if kind == "DOCX":
        if not raw.startswith(b"PK"):
            raise SourceContractError("official DOCX signature is invalid")
        from docx import Document

        document = Document(io.BytesIO(raw))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        blocks.extend(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        return "\n".join(blocks)
    if kind == "DOC":
        if not raw.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
            raise SourceContractError("official DOC signature is invalid")
        converter = shutil.which("textutil")
        if not converter:
            raise SourceContractError("textutil is required for official DOC parsing")
        with tempfile.TemporaryDirectory(prefix="regis-cninfo-doc-") as directory:
            path = Path(directory) / "attachment.doc"
            path.write_bytes(raw)
            result = subprocess.run(
                [converter, "-convert", "txt", "-stdout", str(path)],
                check=False,
                capture_output=True,
                timeout=30,
            )
        if result.returncode != 0:
            raise SourceContractError(
                "official DOC parsing failed: "
                + result.stderr.decode("utf-8", errors="replace")[:240]
            )
        return result.stdout.decode("utf-8", errors="replace")
    raise SourceContractError(f"unsupported official attachment type: {kind}")


def load_historical_official_attachment(
    session: requests.Session, document: dict[str, Any],
) -> tuple[dict[str, Any], str, bytes]:
    """Load and hash one exact original from CNINFO's historical archive."""
    attachment_type = str(document.get("adjunctType") or "").upper()
    attachment_path = str(document.get("adjunctUrl") or "").lstrip("/")
    official_id = str(document.get("announcementId") or "")
    if attachment_type not in {"PDF", "DOCX", "DOC"}:
        raise SourceContractError(
            f"unsupported official attachment type: {attachment_type or 'missing'}"
        )
    if not official_id.isdigit() or not attachment_path:
        raise SourceContractError("official archive attachment identity is invalid")
    url = CNINFO_PDF_BASE + attachment_path
    response = session.get(url, timeout=30)
    response.raise_for_status()
    raw = response.content
    if not raw or len(raw) > 30 * 1024 * 1024:
        raise SourceContractError("official attachment size is invalid")
    text = _extract_original_attachment_text(raw, attachment_type)
    return ({
        "official_id": official_id,
        "official_url": url,
        "official_document_sha256": _sha256_bytes(raw),
        "official_document_bytes": len(raw),
        "official_text_characters": len(text),
        "official_pub_millis": str(document.get("announcementTime") or ""),
        "official_attachment_type": attachment_type,
        "official_source_endpoint": CNINFO_HISTORY_SEARCH_URL,
    }, text, raw)


def load_official_document(
    session: requests.Session, document: dict[str, Any],
) -> tuple[dict[str, Any], str, bytes]:
    url = CNINFO_PDF_BASE + str(document.get("attachmentUrl") or "").lstrip("/")
    response = session.get(url, timeout=30)
    response.raise_for_status()
    raw = response.content
    if not raw.startswith(b"%PDF") or len(raw) > 30 * 1024 * 1024:
        raise SourceContractError("official attachment is not an accepted PDF")
    text = "\n".join(
        page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages
    )
    return ({
        "official_id": str(document.get("attachedId") or ""),
        "official_url": url,
        "official_pdf_sha256": _sha256_bytes(raw),
        "official_pdf_bytes": len(raw),
        "official_text_characters": len(text),
        "official_pub_millis": str(document.get("pubDate") or ""),
    }, text, raw)


def resolve_official_binding(
    session: requests.Session,
    event: dict[str, Any],
) -> dict[str, Any]:
    documents = fetch_official_documents(
        session, event["stock_code"], event["discovery_notice_date"],
    )
    evaluated = []
    for document in documents:
        try:
            evidence, text, _ = load_official_document(session, document)
            evaluated.append({**evidence, **assess_official_text(event, text)})
        except (OSError, ValueError, requests.RequestException, SourceContractError) as exc:
            evaluated.append({
                "official_id": str(document.get("attachedId") or ""),
                "official_url": (
                    CNINFO_PDF_BASE
                    + str(document.get("attachmentUrl") or "").lstrip("/")
                ),
                "valid": False,
                "error": f"{type(exc).__name__}:{exc}",
            })
    valid = [row for row in evaluated if row.get("valid") is True]
    if len(valid) != 1:
        return {
            "status": "unmatched" if not valid else "ambiguous",
            "document_count": len(documents),
            "valid_document_count": len(valid),
            "documents": evaluated,
        }
    selected = valid[0]
    try:
        published_at = dt.datetime.fromtimestamp(
            int(selected["official_pub_millis"]) / 1000, SHANGHAI,
        )
    except (TypeError, ValueError, OSError) as exc:
        return {
            "status": "invalid_official_timestamp",
            "document_count": len(documents),
            "valid_document_count": 1,
            "documents": evaluated,
            "error": f"{type(exc).__name__}:{exc}",
        }
    return {
        "status": "officially_bound",
        "document_count": len(documents),
        "valid_document_count": 1,
        "official_published_at": published_at.isoformat(),
        "official_publication_date": published_at.date().isoformat(),
        "binding": selected,
    }


def _atomic_json(path: Path, payload: dict[str, Any]) -> None:
    encoded = json.dumps(
        payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
        allow_nan=False,
    ).encode()
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(encoded)
            handle.flush()
            os.fsync(handle.fileno())
        Path(temporary).replace(path)
    except Exception:
        Path(temporary).unlink(missing_ok=True)
        raise


def run_pilot(
    *, years: Iterable[int] = range(2022, 2026), output_path: Path = AUDIT_PATH,
) -> dict[str, Any]:
    discovery = requests.Session()
    discovery.headers["User-Agent"] = USER_AGENT
    official = requests.Session()
    official.headers.update({
        "User-Agent": USER_AGENT,
        "Referer": "https://irm.cninfo.com.cn/views/comprehensiveInfo/survey",
        "X-Requested-With": "XMLHttpRequest",
    })
    coverage = []
    sample: list[dict[str, Any]] = []
    for year in years:
        first = fetch_discovery_page(discovery, year, 1)
        pages = int(first.get("pages") or 0)
        page_numbers = sorted({1, max(1, (pages + 1) // 2), max(1, pages)})
        normalized_count = 0
        for page in range(1, pages + 1):
            result = first if page == 1 else fetch_discovery_page(discovery, year, page)
            events = deduplicate_events(result.get("data") or [])
            events = [
                event for event in events
                if (_day(event["discovery_notice_date"]) - _day(
                    event["visit_start_date"]
                )).days <= 16
            ]
            normalized_count += len(events)
            if page in page_numbers:
                sample.extend(events[:2])
        coverage.append({
            "visit_year": year,
            "server_rows": int(first.get("count") or 0),
            "pages": pages,
            "strict_physical_rows_after_calendar_prefilter": normalized_count,
        })

    bindings = []
    for event in sample:
        result = resolve_official_binding(official, event)
        bindings.append({
            "event_id": event["event_id"],
            "stock_code": event["stock_code"],
            "visit_start_date": event["visit_start_date"],
            "discovery_notice_date": event["discovery_notice_date"],
            **result,
        })
    bound = [row for row in bindings if row["status"] == "officially_bound"]
    binding_rate = len(bound) / len(bindings) if bindings else 0.0
    passed = bool(
        len(coverage) >= 4
        and all(row["strict_physical_rows_after_calendar_prefilter"] > 0 for row in coverage)
        and len(bindings) >= 24
        and binding_rate >= 0.90
        and all((row.get("binding") or {}).get("official_pdf_sha256") for row in bound)
    )
    report = {
        "schema": "investor_site_visit_aiv_source_input_audit.v1",
        "candidate_id": "abnormal_investor_site_visits_szse_v1",
        "status": (
            "source_input_feasibility_passed_not_factor"
            if passed else "source_input_feasibility_failed"
        ),
        "audited_at": dt.datetime.now(SHANGHAI).isoformat(timespec="seconds"),
        "primary_source": {
            "title": "Investor Corporate Visits and Predictable Returns",
            "doi": "10.1017/S0022109024000528",
            "official_url": (
                "https://www.cambridge.org/core/journals/"
                "journal-of-financial-and-quantitative-analysis/article/"
                "investor-corporate-visits-and-predictable-returns/"
                "77C958598BB1DC92C46EA32BF252A4B9"
            ),
            "paper_sample": "SZSE July 2012 through December 2019",
            "paper_signal": (
                "cross-sectional residual of log(1 + unique physical visits "
                "published in months m-2..m) on observable visit determinants"
            ),
        },
        "data_contract": {
            "secondary_discovery_only": EASTMONEY_URL,
            "official_timestamp_and_document_authority": CNINFO_SEARCH_URL,
            "official_pdf_base": CNINFO_PDF_BASE,
            "eligible_exchange": "SZSE",
            "require_explicit_physical_visit": True,
            "exclude_remote_or_non_site_activity": list(REMOTE_OR_NON_SITE_TERMS),
            "full_epoch_disclosure_lag_rule": "at most 10 actual trading sessions",
            "factor_time_rule": "official publication timestamp only; never visit date or secondary notice timestamp",
        },
        "coverage": coverage,
        "sample_binding": {
            "method": "first two strict rows from first, middle and last result page in each visit year",
            "sample_size": len(bindings),
            "officially_bound": len(bound),
            "binding_rate": round(binding_rate, 6),
            "required_rate": 0.90,
            "rows": bindings,
        },
        "required_next_steps": [
            "collect and hash-bind the complete 2021-2025 event epoch",
            "replace the 16-calendar-day pilot prefilter with the exact 10-session exchange calendar rule",
            "freeze the float-market-cap SIZE adaptation and SIZE-TURN-MOMEN residual formula before reading returns",
            "run one cost-inclusive validation and sealed out-of-sample test only after input admission",
        ],
        "known_gaps": [
            "the paper uses CSMAR while the free implementation uses Eastmoney only for discovery plus CNINFO for authority",
            "the local panel exposes float market cap rather than the paper's unspecified market-cap convention",
            "historical and production eligibility remain false until a complete official-bound epoch and one-shot tests pass",
        ],
        "outcome_boundary": {
            "stock_returns_read": False,
            "factor_values_computed": False,
            "candidate_ranks_computed": False,
            "ic_computed": False,
            "backtest_computed": False,
            "direction_selected_from_local_outcomes": False,
        },
        "authority": {
            "research_only": True,
            "admission_allowed": False,
            "runnable_factor": False,
            "recommendation_eligible": False,
            "production_strategy": False,
            "generates_buy_or_sell_alert": False,
            "automatic_trading": False,
            "stable_profit_claim": False,
        },
    }
    _atomic_json(output_path, report)
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=AUDIT_PATH)
    args = parser.parse_args(argv)
    report = run_pilot(output_path=args.output)
    print(json.dumps({
        "status": report["status"],
        "coverage": report["coverage"],
        "sample_binding": {
            key: value for key, value in report["sample_binding"].items()
            if key != "rows"
        },
        "output": str(args.output),
        "output_sha256": hashlib.sha256(args.output.read_bytes()).hexdigest(),
    }, ensure_ascii=False))
    return 0 if report["status"] == "source_input_feasibility_passed_not_factor" else 2


if __name__ == "__main__":
    raise SystemExit(main())
